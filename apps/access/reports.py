"""Report queries and exporters (CSV / XLSX)."""

from __future__ import annotations

import csv
from io import BytesIO

from django.http import HttpResponse, StreamingHttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.catalog.models import Application
from apps.orgs.models import Position

from .models import PositionDefault

MATRIX_COLUMNS = [
    "position_code",
    "position",
    "department_code",
    "department",
    "job_code",
    "job_title",
    "position_active",
    "application",
    "access_level",
    "granted_via",
    "target",
    "note",
    "added",
    "added_by",
]


def position_matrix_rows(departments=None, include_inactive=False):
    """One row per (position, default). Positions with no defaults still get a row so
    the export shows gaps."""
    positions = Position.objects.select_related("department", "job_code").order_by("code")
    if departments:
        positions = positions.filter(department__in=departments)
    if not include_inactive:
        positions = positions.filter(is_active=True)
    defaults = (
        PositionDefault.objects.filter(position__in=positions)
        .select_related("access_level__application", "created_by")
        .order_by("position__code", "access_level__application__name", "access_level__name")
    )
    by_position: dict[int, list[PositionDefault]] = {}
    for d in defaults:
        by_position.setdefault(d.position_id, []).append(d)

    for pos in positions:
        base = [
            pos.code,
            pos.display_name,
            pos.department.code,
            pos.department.name,
            pos.job_code.code,
            pos.job_code.title,
            "yes" if pos.is_active else "no",
        ]
        rows = by_position.get(pos.pk)
        if not rows:
            yield base + ["", "", "", "", "", "", ""]
            continue
        for d in rows:
            lvl = d.access_level
            yield base + [
                lvl.application.name,
                lvl.name,
                lvl.get_access_model_display(),
                lvl.access_target,
                d.notes,
                d.created_at.date().isoformat(),
                d.created_by.display_name if d.created_by else "",
            ]


def who_gets(application: Application):
    """{level: [defaults]} for every level of the application (inactive levels included)."""
    levels = list(application.access_levels.order_by("sort_order", "name"))
    defaults = (
        PositionDefault.objects.filter(access_level__application=application)
        .select_related("position__department", "position__job_code", "access_level")
        .order_by("position__code")
    )
    grouped = {lvl: [] for lvl in levels}
    for d in defaults:
        grouped.setdefault(d.access_level, []).append(d)
    return grouped


WHO_GETS_COLUMNS = [
    "application",
    "access_level",
    "granted_via",
    "target",
    "position_code",
    "position",
    "department",
    "job_title",
    "position_active",
    "note",
    "added",
]


def who_gets_rows(application: Application):
    for level, defaults in who_gets(application).items():
        for d in defaults:
            yield [
                application.name,
                level.name,
                level.get_access_model_display(),
                level.access_target,
                d.position.code,
                d.position.display_name,
                d.position.department.name,
                d.position.job_code.title,
                "yes" if d.position.is_active else "no",
                d.notes,
                d.created_at.date().isoformat(),
            ]


# --- Exporters ------------------------------------------------------------------


class _Echo:
    def write(self, value):
        return value


#: What a spreadsheet takes for the start of a formula. Exports carry names and addresses that
#: come from outside the organization -- an Entra ID guest names itself, and any member can name
#: a Microsoft 365 group -- so no cell may reach Excel as one.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _csv_cell(value):
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def csv_response(header, rows, filename) -> StreamingHttpResponse:
    writer = csv.writer(_Echo())

    def stream():
        yield writer.writerow(header)
        for row in rows:
            yield writer.writerow([_csv_cell(value) for value in row])

    response = StreamingHttpResponse(stream(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def xlsx_response(header, rows, filename, sheet_title="Report") -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    widths = [len(h) for h in header]
    for row in rows:
        ws.append(row)
        for i, (cell, value) in enumerate(zip(ws[ws.max_row], row, strict=False)):
            if isinstance(value, str) and value.startswith("="):
                # openpyxl stores a string starting with "=" as a formula; keep it text.
                cell.data_type = "s"
            widths[i] = min(max(widths[i], len(str(value or ""))), 60)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w + 2
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
