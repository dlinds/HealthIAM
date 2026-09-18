"""Browser-facing pieces of the Active Directory integration: the group picker, the catalog
hooks (reference badges, tab count, form attributes), the AD groups page, the broken-reference
report, Admin > Active Directory and the nav / dashboard / Users page entries."""

import io
from datetime import timedelta

import pytest
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.catalog.models import AccessLevel
from apps.directory import references
from apps.directory.models import ADGroup, DirectorySyncRun

from . import factories

SECRET = "test-secret-not-real"  # config/settings/test.py AD_BIND_PASSWORD

pytestmark = pytest.mark.django_db

PICKER_ATTRS = (
    b'hx-target="#directory-group-picker"',
    b'hx-swap="innerHTML"',
    b'hx-trigger="focus once, input changed delay:250ms"',
    b'autocomplete="off"',
)


@pytest.fixture
def app(db):
    return factories.ApplicationFactory(name="Epic")


@pytest.fixture
def analyst_user(db, app):
    user = factories.UserFactory(username="analyst")
    factories.make_analyst(app, user, is_primary=True)
    return user


def groups_synced():
    return DirectorySyncRun.objects.create(scope="groups", status="completed")


# --- Group picker -----------------------------------------------------------------------------


def test_picker_reads_the_input_name_marks_verified_and_renders_buttons(as_user, help_desk_user):
    factories.ADGroupFactory(name="APP_PACS_VIEW", description="PACS viewers")
    factories.ADGroupFactory(name="APP_PACS_VIEW_EXT")
    factories.ADGroupFactory(name="APP_EPIC_RN", cn="Epic nurses", description="Epic nurse pool")
    gone = factories.ADGroupFactory(name="APP_PACS_OLD")
    gone.deactivate()
    client = as_user(help_desk_user)
    url = reverse("directory:group_picker")

    resp = client.get(url, {"ad_group_name": "app_pacs_view"})
    assert resp.status_code == 200
    body = resp.content.decode()
    assert body.count('type="button"') == 2
    assert 'data-pick-group="APP_PACS_VIEW"' in body
    assert 'data-pick-group="APP_PACS_VIEW_EXT"' in body
    assert "APP_PACS_OLD" not in body and "APP_EPIC_RN" not in body
    assert body.count("Verified") == 1
    assert body.index("APP_PACS_VIEW") < body.index("APP_PACS_VIEW_EXT")
    assert "<input" not in body and "<a " not in body
    assert "PACS viewers" in body

    # `q` is an alias; description and cn are searched too; no exact match -> no "Verified"
    resp = client.get(url, {"q": "nurse"})
    body = resp.content.decode()
    assert 'data-pick-group="APP_EPIC_RN"' in body and "Verified" not in body
    resp = client.get(url, {"q": "epic nurses"})
    assert b"APP_EPIC_RN" in resp.content

    resp = client.get(url, {"ad_group_name": "nothing-like-this"})
    assert resp.status_code == 200
    assert b"No imported group matches. Free text is fine for groups outside the sync filter." in (
        resp.content
    )
    assert b'type="button"' not in resp.content


def test_picker_limits_to_fifteen_active_groups(as_user, help_desk_user):
    for n in range(20):
        factories.ADGroupFactory(name=f"APP_MANY_{n:02d}")
    client = as_user(help_desk_user)
    resp = client.get(reverse("directory:group_picker"), {"ad_group_name": "APP_MANY"})
    assert resp.content.decode().count("data-pick-group=") == 15
    resp = client.get(reverse("directory:group_picker"))  # focus once: empty query lists groups
    assert resp.status_code == 200
    assert resp.content.decode().count("data-pick-group=") == 15


def test_picker_requires_a_role(client, plain_user):
    client.force_login(plain_user)
    assert client.get(reverse("directory:group_picker"), {"q": "x"}).status_code == 403
    client.logout()
    resp = client.get(reverse("directory:group_picker"), {"q": "x"})
    assert resp.status_code == 302
    assert resp.url.startswith(reverse("accounts:login"))


# --- Access-level form and tab -----------------------------------------------------------


def test_access_level_form_carries_picker_attributes_and_free_text_saves(
    as_user, analyst_user, app
):
    client = as_user(analyst_user)
    url = reverse("catalog:access_level_add", args=[app.pk])
    resp = client.get(url, HTTP_HX_REQUEST="true")
    assert resp.status_code == 200
    body = resp.content
    for attr in PICKER_ATTRS:
        assert attr in body
    picker_url = reverse("directory:group_picker").encode()
    assert b'hx-get="' + picker_url + b'"' in body
    assert b'<div id="directory-group-picker" class="list-group list-group-flush small mt-1">' in (
        body
    )
    assert b"[data-pick-group]" in body and b'new Event("input", { bubbles: true })' in body

    # Same body as tests/test_catalog.py: the picker never becomes a hard constraint
    resp = client.post(
        url,
        {
            "name": "Nurse",
            "access_model": "ad_group",
            "ad_group_name": "APP_EPIC_NURSE",
            "sort_order": 10,
            "is_active": "on",
        },
    )
    assert resp.status_code == 200 and "HX-Retarget" not in resp.headers
    assert AccessLevel.objects.get(application=app, name="Nurse").ad_group_name == "APP_EPIC_NURSE"

    resp = client.post(
        url,
        {
            "name": "Custom",
            "access_model": "ad_group",
            "ad_group_name": "SG-Custom Team",
            "sort_order": 20,
            "is_active": "on",
        },
    )
    assert resp.status_code == 200 and "HX-Retarget" not in resp.headers
    assert AccessLevel.objects.get(application=app, name="Custom").ad_group_name == "SG-Custom Team"


@override_settings(AD_ENABLED=False)
def test_access_level_form_is_plain_when_ad_is_disabled(as_user, analyst_user, app):
    client = as_user(analyst_user)
    resp = client.get(reverse("catalog:access_level_add", args=[app.pk]), HTTP_HX_REQUEST="true")
    assert resp.status_code == 200
    for attr in PICKER_ATTRS:
        assert attr not in resp.content
    assert b'<div id="directory-group-picker"' not in resp.content
    assert b"Start typing to pick an imported AD group" not in resp.content


def test_access_levels_tab_shows_badges_and_broken_count(as_user, help_desk_user, app):
    groups_synced()
    factories.ADGroupFactory(name="APP_EPIC_RN")
    old = factories.ADGroupFactory(name="APP_EPIC_OLD")
    old.deactivate()
    factories.AccessLevelFactory(application=app, name="Nurse", ad_group_name="app_epic_rn")
    factories.AccessLevelFactory(application=app, name="Old", ad_group_name="APP_EPIC_OLD")
    factories.AccessLevelFactory(application=app, name="Gone", ad_group_name="APP_EPIC_GONE")
    factories.AccessLevelFactory(application=app, name="Custom", ad_group_name="SG-Custom")
    factories.AccessLevelFactory(
        application=app,
        name="Desk",
        access_model="ticket",
        ad_group_name="",
        ticket_assignment_team="SD",
    )
    client = as_user(help_desk_user)
    resp = client.get(reverse("catalog:application_detail", args=[app.pk]))
    assert resp.status_code == 200
    body = resp.content.decode()
    assert body.count(">In AD<") == 1
    assert "Not returned by the last sync (last seen " in body
    assert body.count("Not found in AD") == 1
    assert body.count("outside sync filter") == 1
    assert 'Access levels <span class="badge text-bg-secondary">5</span>' in body
    assert resp.context["broken_level_count"] == 2
    assert "2</span></button>" in body  # warning count on the tab label
    assert "text-bg-warning" in body.split('data-bs-target="#tab-levels"')[1].split("</button>")[0]


def test_access_levels_tab_has_no_count_or_badges_before_the_first_sync(
    as_user, help_desk_user, app
):
    factories.AccessLevelFactory(application=app, name="Nurse", ad_group_name="APP_EPIC_RN")
    client = as_user(help_desk_user)
    resp = client.get(reverse("catalog:application_detail", args=[app.pk]))
    body = resp.content.decode()
    assert resp.context["broken_level_count"] == 0
    texts = ("In AD", "Not found in AD", "Not returned by the last sync", "outside sync filter")
    for text in texts:
        assert text not in body
    tab = body.split('data-bs-target="#tab-levels"')[1].split("</button>")[0]
    assert "text-bg-warning" not in tab


def test_htmx_section_after_saving_a_level_carries_the_badge(as_user, analyst_user, app):
    groups_synced()
    factories.ADGroupFactory(name="APP_EPIC_NURSE")
    client = as_user(analyst_user)
    resp = client.post(
        reverse("catalog:access_level_add", args=[app.pk]),
        {
            "name": "Nurse",
            "access_model": "ad_group",
            "ad_group_name": "APP_EPIC_NURSE",
            "sort_order": 10,
            "is_active": "on",
        },
    )
    assert resp.status_code == 200 and b">In AD<" in resp.content


# --- AD groups page -----------------------------------------------------------------------


def test_group_list_filters_and_references_for_help_desk(as_user, help_desk_user, app):
    groups_synced()
    pacs = factories.ADGroupFactory(name="APP_PACS_VIEW", description="PACS viewers")
    factories.ADGroupFactory(name="APP_EPIC_RN", cn="Epic nurses", category="distribution")
    factories.ADGroupFactory(name="LIC_M365_E3", description="Microsoft 365 licence")
    gone = factories.ADGroupFactory(name="APP_PACS_OLD")
    gone.deactivate()
    level = factories.AccessLevelFactory(application=app, name="Nurse", ad_group_name="app_epic_rn")
    client = as_user(help_desk_user)
    url = reverse("directory:group_list")

    resp = client.get(url)
    assert resp.status_code == 200
    names = [g.name for g in resp.context["object_list"]]
    assert names == ["APP_EPIC_RN", "APP_PACS_VIEW", "LIC_M365_E3"]  # inactive hidden by default
    body = resp.content.decode()
    assert "Epic · Nurse" in body and level.get_absolute_url() in body
    assert reverse("directory:admin_index") not in body  # admin-only link
    epic = next(g for g in resp.context["object_list"] if g.name == "APP_EPIC_RN")
    assert [lvl.pk for lvl in epic.referencing_levels] == [level.pk]
    assert next(g for g in resp.context["object_list"] if g.pk == pacs.pk).referencing_levels == []

    resp = client.get(url, {"q": "pacs", "active": "all"})
    assert [g.name for g in resp.context["object_list"]] == ["APP_PACS_OLD", "APP_PACS_VIEW"]
    resp = client.get(url, {"q": "epic nurses"})  # cn
    assert [g.name for g in resp.context["object_list"]] == ["APP_EPIC_RN"]
    resp = client.get(url, {"q": "365"})  # description
    assert [g.name for g in resp.context["object_list"]] == ["LIC_M365_E3"]
    resp = client.get(url, {"q": "OU=Groups", "active": "0"})  # DN
    assert [g.name for g in resp.context["object_list"]] == ["APP_PACS_OLD"]
    resp = client.get(url, {"category": "distribution"})
    assert [g.name for g in resp.context["object_list"]] == ["APP_EPIC_RN"]
    resp = client.get(url, {"unreferenced": "1"})
    assert [g.name for g in resp.context["object_list"]] == ["APP_PACS_VIEW", "LIC_M365_E3"]
    assert resp.context["unreferenced"] is True


def test_group_list_shows_admin_link_and_unsynced_note_for_admin(as_user, admin_user):
    resp = as_user(admin_user).get(reverse("directory:group_list"))
    assert resp.status_code == 200
    body = resp.content.decode()
    assert reverse("directory:admin_index") in body
    assert "No group sync has completed yet." in body
    assert "No AD groups found." in body


# --- Broken-reference report ----------------------------------------------------------------


@pytest.fixture
def broken_world(db, app):
    groups_synced()
    factories.ADGroupFactory(name="APP_EPIC_RN")
    old = factories.ADGroupFactory(name="APP_EPIC_OLD")
    old.deactivate()
    ok = factories.AccessLevelFactory(application=app, name="Nurse", ad_group_name="APP_EPIC_RN")
    inactive = factories.AccessLevelFactory(
        application=app, name="Old", ad_group_name="APP_EPIC_OLD"
    )
    missing = factories.AccessLevelFactory(
        application=app, name="Gone", ad_group_name="APP_EPIC_GONE"
    )
    unverified = factories.AccessLevelFactory(
        application=app, name="Custom", ad_group_name="SG-Custom"
    )
    return {"ok": ok, "inactive": inactive, "missing": missing, "unverified": unverified}


def test_broken_reference_page_csv_and_xlsx(as_user, help_desk_user, broken_world):
    client = as_user(help_desk_user)
    url = reverse("directory:broken_references")

    resp = client.get(url)
    assert resp.status_code == 200
    body = resp.content.decode()
    assert [row["level"].name for row in resp.context["rows"]] == ["Gone", "Old"]
    assert "APP_EPIC_GONE" in body and "Not found in AD" in body
    assert "APP_EPIC_OLD" in body and "Not returned by the last sync" in body
    assert "SG-Custom" not in body and "APP_EPIC_RN" not in body
    assert "?format=csv" in body and "?format=xlsx" in body

    resp = client.get(url, {"format": "csv"})
    assert resp["Content-Type"].startswith("text/csv")
    assert resp["Content-Disposition"] == 'attachment; filename="broken-ad-references.csv"'
    text = b"".join(resp.streaming_content).decode()
    lines = text.splitlines()
    assert lines[0] == ",".join(references.BROKEN_REF_COLUMNS)
    assert len(lines) == 3
    assert "APP_EPIC_GONE" in text and "APP_EPIC_OLD" in text and "SG-Custom" not in text

    resp = client.get(url, {"format": "xlsx"})
    assert resp["Content-Type"].endswith("spreadsheetml.sheet")
    assert resp["Content-Disposition"] == 'attachment; filename="broken-ad-references.xlsx"'
    ws = load_workbook(io.BytesIO(resp.content)).active
    header = [cell.value for cell in ws[1]]
    assert header == references.BROKEN_REF_COLUMNS
    assert ws.max_row == 3
    names = {ws.cell(row=r, column=5).value for r in (2, 3)}
    assert names == {"APP_EPIC_GONE", "APP_EPIC_OLD"}


def test_broken_reference_page_before_first_sync(as_user, help_desk_user, app):
    factories.AccessLevelFactory(application=app, name="Nurse", ad_group_name="APP_EPIC_RN")
    resp = as_user(help_desk_user).get(reverse("directory:broken_references"))
    assert resp.status_code == 200
    assert resp.context["rows"] == []
    assert b"has not been imported yet" in resp.content


# --- Permission matrix ----------------------------------------------------------------------

READ_ONLY = ("group_list", "group_picker", "broken_references")
ADMIN_GET = ("admin_index", "run_list", "run_detail")
ADMIN_POST = ("connection_test", "sync_start", "run_apply")


def _url(name, run):
    if name in ("run_detail", "run_apply"):
        return reverse(f"directory:{name}", args=[run.pk])
    return reverse(f"directory:{name}")


def test_permission_matrix_over_all_directory_urls(
    as_user, help_desk_user, admin_user, plain_user, client, fake_directory
):
    run = DirectorySyncRun.objects.create(scope="groups", status="completed")
    assert len(READ_ONLY) + len(ADMIN_GET) + len(ADMIN_POST) == 9

    c = as_user(help_desk_user)
    for name in READ_ONLY:
        assert c.get(_url(name, run)).status_code == 200, name
    for name in ADMIN_GET:
        assert c.get(_url(name, run)).status_code == 403, name
    for name in ADMIN_POST:
        assert c.get(_url(name, run)).status_code == 405, name
        assert c.post(_url(name, run), {"scope": "groups"}).status_code == 403, name
    assert DirectorySyncRun.objects.count() == 1  # help desk started nothing

    c = as_user(admin_user)
    for name in READ_ONLY + ADMIN_GET:
        assert c.get(_url(name, run)).status_code == 200, name
    for name in ADMIN_POST:
        assert c.get(_url(name, run)).status_code == 405, name

    c = as_user(plain_user)
    for name in READ_ONLY + ADMIN_GET:
        assert c.get(_url(name, run)).status_code == 403, name

    client.logout()
    for name in READ_ONLY + ADMIN_GET:
        resp = client.get(_url(name, run))
        assert resp.status_code == 302 and resp.url.startswith(reverse("accounts:login")), name


# --- Admin > Active Directory ---------------------------------------------------------------


def test_admin_index_shows_config_status_schedule_and_runs_without_the_secret(
    as_user, admin_user, fake_directory
):
    groups_synced()
    factories.ADGroupFactory(name="APP_EPIC_RN")
    gone = factories.ADGroupFactory(name="APP_EPIC_OLD")
    gone.deactivate()
    factories.UserFactory(username="alice", ad_managed=True)
    factories.UserFactory(username="carol", ad_managed=True, is_active=False)
    factories.AccessLevelFactory(name="Gone", ad_group_name="APP_EPIC_GONE")
    failed = DirectorySyncRun.objects.create(
        scope="all", status="failed", error=f"bind failed for {SECRET}", created_by=admin_user
    )
    resp = as_user(admin_user).get(reverse("directory:admin_index"))
    assert resp.status_code == 200
    body = resp.content.decode()
    assert SECRET not in body
    config = resp.context["config"]
    assert "bind_password" not in config and config["bind_password_set"] is True
    assert "ldaps://dc.test.invalid" in body and "DC=test,DC=invalid" in body
    assert "OU=Groups,DC=test,DC=invalid" in body and "APP_*, LIC_*" in body
    assert "docker exec ix-healthiam-web-1 python manage.py sync_ad" in body
    assert resp.context["groups_active"] == 1 and resp.context["groups_inactive"] == 1
    assert resp.context["managed_active"] == 1 and resp.context["managed_inactive"] == 1
    assert resp.context["broken_count"] == 1
    assert resp.context["last_run"] == failed
    assert reverse("directory:run_detail", args=[failed.pk]) in body
    assert 'hx-post="' + reverse("directory:connection_test") + '"' in body
    assert 'hx-target="#connection-result"' in body and 'id="connection-result"' in body
    assert 'hx-indicator="#sync-indicator"' in body and "hx-disabled-elt" in body
    assert '<select name="scope" class="form-select"' in body or (
        '<select name="scope" id="id_scope" class="form-select"' in body
    )
    assert resp.context["check_warnings"] == []  # test settings are clean


def test_admin_index_before_any_sync(as_user, admin_user):
    resp = as_user(admin_user).get(reverse("directory:admin_index"))
    assert resp.status_code == 200
    assert resp.context["last_run"] is None and resp.context["broken_count"] is None
    assert b"not checked yet" in resp.content and b"No sync runs yet." in resp.content


def test_admin_index_shows_check_warnings_inline(as_user, admin_user, settings):
    settings.AD_BIND_PASSWORD = ""
    resp = as_user(admin_user).get(reverse("directory:admin_index"))
    assert [w.id for w in resp.context["check_warnings"]] == ["directory.W003"]
    assert b"directory.W003" in resp.content


def test_connection_test_ok_and_failure_are_200_and_close_the_client(
    as_user, admin_user, fake_directory
):
    client = as_user(admin_user)
    url = reverse("directory:connection_test")

    resp = client.post(url)
    assert resp.status_code == 200
    body = resp.content.decode()
    assert "<html" not in body and "Connected" in body and "alert-success" in body
    assert "fake-dc.test.invalid" in body
    assert "CN=IAM-Users,OU=IAM,DC=test,DC=invalid" in body
    assert fake_directory.closed is True

    fake_directory.closed = False
    fake_directory.fail_connect = f"LDAPS handshake failed; bind with {SECRET}"
    resp = client.post(url)
    assert resp.status_code == 200
    body = resp.content.decode()
    assert "Connection failed" in body and "alert-danger" in body
    assert "LDAPS handshake failed" in body and SECRET not in body
    assert fake_directory.closed is True


def test_connection_test_survives_a_client_construction_error(as_user, admin_user, monkeypatch):
    def boom():
        raise RuntimeError(f"cannot build client with {SECRET}")

    monkeypatch.setattr("apps.directory.sync.build_client", boom)
    resp = as_user(admin_user).post(reverse("directory:connection_test"))
    assert resp.status_code == 200
    body = resp.content.decode()
    assert "Connection failed" in body and "RuntimeError: cannot build client" in body
    assert SECRET not in body


def test_sync_now_preview_then_apply_on_the_same_run(as_user, admin_user, fake_directory):
    client = as_user(admin_user)
    resp = client.post(reverse("directory:sync_start"), {"scope": "all"})
    run = DirectorySyncRun.objects.get()
    assert resp.status_code == 302 and resp.url == run.get_absolute_url()
    assert run.status == DirectorySyncRun.Status.PREVIEWED
    assert run.trigger == DirectorySyncRun.Trigger.MANUAL and run.created_by == admin_user
    assert run.server == "fake-dc.test.invalid" and fake_directory.closed is True
    assert run.summary["users"]["created"] == 3 and run.summary["groups"]["created"] == 3
    # Dry run: nothing written.
    assert not User.objects.filter(username="alice@test.invalid").exists()
    assert ADGroup.objects.count() == 0

    resp = client.get(run.get_absolute_url())
    assert resp.status_code == 200
    body = resp.content.decode()
    assert "Apply sync" in body and reverse("directory:run_apply", args=[run.pk]) in body
    assert "confirm(" in body and "This is a dry run" in body
    assert body.count("<th>Kind</th>") == 1  # Changes table (no Problems table on a clean run)
    assert "alice@test.invalid" in body and "APP_PACS_VIEW" in body
    assert [(kind, part["created"]) for kind, part in resp.context["parts"]] == [
        ("users", 3),
        ("groups", 3),
    ]
    assert "Skipped" in body

    resp = client.post(reverse("directory:run_apply", args=[run.pk]), follow=True)
    assert resp.status_code == 200
    run.refresh_from_db()
    assert run.status == DirectorySyncRun.Status.COMPLETED
    assert DirectorySyncRun.objects.count() == 1
    assert User.objects.get(username="alice@test.invalid").ad_managed is True
    assert set(ADGroup.objects.values_list("name", flat=True)) == {
        "APP_PACS_VIEW",
        "APP_EPIC_RN",
        "LIC_M365_E3",
    }
    text = [str(m) for m in resp.context["messages"]]
    assert text == [
        "Sync applied — users: 3 created, 0 updated, 0 reactivated, 0 deactivated, 0 errors; "
        "groups: 3 created, 0 updated, 0 reactivated, 0 deactivated, 0 errors."
    ]
    assert b"Apply sync" not in resp.content and b"Completed" in resp.content

    # A completed run cannot be applied again.
    before = User.objects.count()
    resp = client.post(reverse("directory:run_apply", args=[run.pk]), follow=True)
    assert [str(m) for m in resp.context["messages"]] == ["Only a previewed sync can be applied."]
    run.refresh_from_db()
    assert run.status == DirectorySyncRun.Status.COMPLETED and User.objects.count() == before


def test_sync_now_scope_and_htmx_redirect(as_user, admin_user, fake_directory):
    client = as_user(admin_user)
    resp = client.post(reverse("directory:sync_start"), {"scope": "groups"}, HTTP_HX_REQUEST="true")
    run = DirectorySyncRun.objects.get()
    assert resp.status_code == 200 and resp["HX-Redirect"] == run.get_absolute_url()
    assert run.scope == DirectorySyncRun.Scope.GROUPS
    assert run.summary["users"] is None and run.summary["groups"]["created"] == 3

    resp = client.post(reverse("directory:sync_start"), {"scope": "bogus"}, follow=True)
    assert DirectorySyncRun.objects.count() == 1
    assert [str(m) for m in resp.context["messages"]] == ["Choose what to sync."]
    assert resp.redirect_chain[-1][0] == reverse("directory:admin_index")


def test_sync_now_failure_is_a_failed_run_not_a_500(as_user, admin_user, fake_directory):
    fake_directory.fail_connect = f"bind failed for {SECRET}"
    client = as_user(admin_user)
    resp = client.post(reverse("directory:sync_start"), {"scope": "all"}, follow=True)
    assert resp.status_code == 200
    run = DirectorySyncRun.objects.get()
    assert run.status == DirectorySyncRun.Status.FAILED
    assert run.error.startswith("DirectoryUnavailable: bind failed for") and SECRET not in run.error
    body = resp.content.decode()
    assert SECRET not in body
    assert "The run failed; nothing was written." in body and "Failed" in body
    assert "Apply sync" not in body
    messages = [str(m) for m in resp.context["messages"]]
    assert len(messages) == 1 and messages[0].startswith("Sync preview failed: ")
    assert SECRET not in messages[0]

    resp = client.post(reverse("directory:run_apply", args=[run.pk]), follow=True)
    assert [str(m) for m in resp.context["messages"]] == ["Only a previewed sync can be applied."]

    resp = client.get(reverse("directory:run_list"))
    assert resp.status_code == 200 and SECRET not in resp.content.decode()
    assert b">Failed<" in resp.content


def _entry(row, code, action, message):
    entry = {"kind": "users", "row": row, "code": code, "action": action}
    return {**entry, "message": message, "dn": ""}


def test_run_detail_shows_problems_stale_and_skipped(as_user, admin_user):
    run = DirectorySyncRun.objects.create(
        scope="users",
        status="previewed",
        summary={
            "users": {
                "created": 0,
                "updated": 1,
                "reactivated": 0,
                "deactivated": 0,
                "unchanged": 2,
                "errors": 1,
                "rows": 4,
                "skipped": 1,
            },
            "groups": None,
        },
        log=[
            _entry(2, "x@t", "error", "No UPN"),
            _entry(3, "y@t", "updated", "title"),
            _entry(0, "IAM-Users", "skipped", "Missing-member pass skipped"),
        ],
    )
    resp = as_user(admin_user).get(run.get_absolute_url())
    body = resp.content.decode()
    assert "Entries with errors (1)" in body and "No UPN" in body
    assert "Changes (2)" in body and "Missing-member pass skipped" in body
    assert body.count("<th>Kind</th>") == 2
    assert [e["code"] for e in resp.context["problems"]] == ["x@t"]
    assert [e["action"] for e in resp.context["changes"]] == ["updated", "skipped"]

    stale = DirectorySyncRun.objects.create(
        scope="all", status="pending", started_at=timezone.now() - timedelta(minutes=30)
    )
    resp = as_user(admin_user).get(stale.get_absolute_url())
    assert b"Abandoned (worker stopped)" in resp.content and b"Apply sync" not in resp.content
    resp = as_user(admin_user).get(reverse("directory:run_list"))
    assert b"Abandoned (worker stopped)" in resp.content


# --- Dashboard, nav, report card, AD_ENABLED=False ------------------------------------------


def test_dashboard_shows_the_broken_reference_quality_entry(as_user, admin_user, broken_world):
    resp = as_user(admin_user).get(reverse("core:dashboard"))
    assert resp.status_code == 200
    count, sample = resp.context["quality"]["broken_ad_references"]
    assert count == 2
    assert [lvl.name for lvl in sample] == ["Gone", "Old"]
    body = resp.content.decode()
    assert "Access levels pointing at AD groups not found" in body
    assert "Epic · Gone" in body and "Epic · Old" in body
    assert broken_world["missing"].get_absolute_url() in body
    assert broken_world["missing"].get_absolute_url().endswith("#tab-levels")


def test_nav_and_report_card_show_directory_entries(as_user, admin_user, help_desk_user):
    resp = as_user(admin_user).get(reverse("access:reports_index"))
    body = resp.content.decode()
    assert "Broken AD references" in body
    assert reverse("directory:broken_references") + "?format=csv" in body
    assert reverse("directory:broken_references") + "?format=xlsx" in body
    assert ">AD groups</a>" in body and ">Active Directory</a>" in body

    body = as_user(help_desk_user).get(reverse("access:reports_index")).content.decode()
    assert ">AD groups</a>" in body and ">Active Directory</a>" not in body


@override_settings(AD_ENABLED=False)
def test_ad_disabled_hides_nav_report_card_and_dashboard_entry(as_user, admin_user, broken_world):
    client = as_user(admin_user)
    resp = client.get(reverse("core:dashboard"))
    assert resp.status_code == 200
    assert "broken_ad_references" not in resp.context["quality"]
    body = resp.content.decode()
    assert "Access levels pointing at AD groups not found" not in body
    assert ">AD groups</a>" not in body and ">Active Directory</a>" not in body

    body = client.get(reverse("access:reports_index")).content.decode()
    assert "Broken AD references" not in body and "broken-ad-references" not in body
    assert reverse("directory:broken_references") not in body


# --- Users page ------------------------------------------------------------------------------


def test_users_page_marks_ad_managed_logins_and_explains_the_sync(as_user, admin_user):
    managed = factories.UserFactory(username="alice@test.invalid", ad_managed=True)
    plain = factories.UserFactory(username="bob")
    client = as_user(admin_user)

    resp = client.get(reverse("accounts:user_list"))
    body = resp.content.decode()

    def row_for(u):
        return body.split(u.username)[0].rsplit("<tr>", 1)[1]

    assert ">AD</span>" in row_for(managed) and ">AD</span>" not in row_for(plain)
    assert "or by the Active Directory sync" in body

    resp = client.get(reverse("accounts:user_roles", args=[managed.pk]))
    body = resp.content.decode()
    assert "alert-info" in body and "Active Directory sync" in body
    assert "Account active" in body and "IAM-Users" in body and "Help Desk" in body
    resp = client.get(reverse("accounts:user_roles", args=[plain.pk]))
    assert "alert-info" not in resp.content.decode()


@override_settings(AD_ENABLED=False)
def test_users_footer_is_unchanged_when_ad_is_disabled(as_user, admin_user):
    body = as_user(admin_user).get(reverse("accounts:user_list")).content.decode()
    assert "Users are created on first SSO sign-in." in body
    assert "Active Directory sync" not in body
