import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.access import reports, services

from . import factories

pytestmark = pytest.mark.django_db


@pytest.fixture
def world(db, admin_user):
    epic = factories.ApplicationFactory(name="Epic")
    nurse = factories.AccessLevelFactory(
        application=epic, name="Nurse", ad_group_name="APP_EPIC_RN"
    )
    clerk = factories.AccessLevelFactory(
        application=epic, name="Clerk", ad_group_name="APP_EPIC_CLERK"
    )
    nursing = factories.DepartmentFactory(code="0100", name="Nursing")
    admin_dept = factories.DepartmentFactory(code="0200", name="Administration")
    rn = factories.PositionFactory(
        department=nursing, job_code=factories.JobCodeFactory(code="7000", title="RN")
    )
    clerk_pos = factories.PositionFactory(
        department=admin_dept, job_code=factories.JobCodeFactory(code="8000", title="Unit Clerk")
    )
    empty = factories.PositionFactory(
        department=nursing, job_code=factories.JobCodeFactory(code="7001", title="LPN")
    )
    services.add_default(rn, nurse, actor=admin_user, reason="seed", notes="all units")
    services.add_default(rn, clerk, actor=admin_user, reason="seed")
    services.add_default(clerk_pos, clerk, actor=admin_user, reason="seed")
    return {
        "epic": epic,
        "nurse": nurse,
        "clerk": clerk,
        "rn": rn,
        "clerk_pos": clerk_pos,
        "empty": empty,
        "nursing": nursing,
    }


def test_matrix_rows_include_positions_without_defaults(world):
    rows = list(reports.position_matrix_rows())
    codes = [(r[0], r[8]) for r in rows]
    assert ("0100-7000", "Nurse") in codes
    assert ("0100-7000", "Clerk") in codes
    assert ("0200-8000", "Clerk") in codes
    assert ("0100-7001", "") in codes
    nursing_rows = list(reports.position_matrix_rows(departments=[world["nursing"]]))
    assert all(r[2] == "0100" for r in nursing_rows) and len(nursing_rows) == 3


def test_matrix_export_csv_and_xlsx(as_user, help_desk_user, world):
    client = as_user(help_desk_user)
    resp = client.get(
        reverse("access:reports_index"),
        {"report": "matrix", "format": "csv", "departments": [world["nursing"].pk]},
    )
    assert resp["Content-Type"].startswith("text/csv")
    body = b"".join(resp.streaming_content).decode()
    assert body.splitlines()[0].startswith("position_code,position,department_code")
    assert "APP_EPIC_RN" in body and "0200-8000" not in body

    resp = client.get(reverse("access:reports_index"), {"report": "matrix", "format": "xlsx"})
    assert resp["Content-Type"].endswith("spreadsheetml.sheet")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws["A1"].value == "position_code"
    assert ws.max_row == 1 + 4  # 3 defaults + 1 empty position


def test_who_gets_page_and_export(as_user, help_desk_user, world):
    client = as_user(help_desk_user)
    url = reverse("access:who_gets", args=[world["epic"].pk])
    resp = client.get(url)
    assert resp.status_code == 200
    assert resp.context["total"] == 3
    assert b"0100-7000" in resp.content and b"0200-8000" in resp.content

    resp = client.get(url, {"format": "csv"})
    body = b"".join(resp.streaming_content).decode()
    assert body.count("0100-7000") == 2 and "Clerk" in body

    resp = client.get(reverse("access:reports_index"))
    assert resp.status_code == 200 and b"Position access matrix" in resp.content


def test_reports_need_a_role(as_user, plain_user):
    assert as_user(plain_user).get(reverse("access:reports_index")).status_code == 403
